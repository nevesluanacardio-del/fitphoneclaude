#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera a planilha "Diagnóstico do Gargalo Oculto" (.xlsx) pronta para subir no
Google Sheets: menus suspensos por item, pontuação automática, gráfico radar
e leitura do gargalo (dimensão de menor pontuação).

Pontuação (igual ao app):
  "Existe e é satisfatório"        -> peso cheio
  "Existe, mas não é satisfatório" -> metade do peso
  "Não existe"                     -> 0
Cada dimensão vale no máximo 100 pontos (= % no radar).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import RadarChart, Reference
from openpyxl.utils import get_column_letter

OPT_FULL = "Existe e é satisfatório"
OPT_PARTIAL = "Existe, mas não é satisfatório"
OPT_NONE = "Não existe"

DIMENSIONS = [
    {
        "id": "processos", "name": "Processos", "weight": 10, "color": "3B82F6",
        "intro": "Como o trabalho realmente acontece no dia a dia.",
        "questions": [
            "Os principais processos do negócio estão documentados.",
            "Quando aparecem problemas, não somente apagamos incêndios — as causas são investigadas.",
            "Os colaboradores têm clareza da missão, visão, valores, políticas, ações e objetivos.",
            "Existem indicadores mapeados e monitorados.",
            "Acontecem reuniões de alinhamento das tarefas com as equipes (gerenciamento da rotina).",
            "Os colaboradores sabem exatamente o que fazer, pois existe um manual (ou documento) de cultura/orientações.",
            "São realizados planos de ação para alcançar as metas/objetivos.",
            "Existe uma dose adequada de planejamento antes da execução.",
            "Temos uma mentalidade de melhoria contínua, utilizando a ferramenta PDCA.",
            "Temos um organograma documentado e compartilhado com todos.",
        ],
        "titulo": "Seu gargalo oculto está nos PROCESSOS",
        "diagnostico": ("Seu negócio cresceu, mas a operação ainda depende de você e de alguns \"heróis\". "
                        "Sem processos documentados e rotina gerenciada, cada novo cliente adiciona caos em vez de lucro. "
                        "O faturamento parou de crescer porque a empresa não consegue entregar mais sem quebrar."),
    },
    {
        "id": "resultados", "name": "Resultados", "weight": 20, "color": "10B981",
        "intro": "O que o negócio efetivamente entrega e mede.",
        "questions": [
            "A margem de lucro do negócio é mensurada e está em nível satisfatório.",
            "A satisfação do cliente é mensurada e está em nível satisfatório.",
            "A satisfação do colaborador é mensurada e está em nível satisfatório.",
            "Existe remuneração atrelada ao desempenho (ao menos para as pessoas-chave) na organização.",
            "A taxa de crescimento da empresa geralmente é maior que a esperada/planejada.",
        ],
        "titulo": "Seu gargalo oculto está nos RESULTADOS",
        "diagnostico": ("Você está ocupado o tempo todo, mas não sabe ao certo se está lucrando. "
                        "Sem medir margem, satisfação de cliente e de equipe, você pilota a empresa no escuro. "
                        "O faturamento estagnou porque você otimiza o que sente, não o que os números mostram."),
    },
    {
        "id": "gestao", "name": "Sistema de Gestão", "weight": 20, "color": "F59E0B",
        "intro": "O leme estratégico que conduz a empresa.",
        "questions": [
            "É feito planejamento anualmente para curto, médio e longo prazo (1 a 3 ou 1 a 5 anos).",
            "Os sistemas de informação (softwares) são adequados para o nível de maturidade/momento da empresa.",
            "As metas são bem estabelecidas, sempre contemplando: indicador, prazo e valor.",
            "Existe por parte dos donos/líderes alocação de tempo adequado para a parte estratégica da empresa.",
            "Os indicadores são acompanhados e comunicados para a equipe.",
        ],
        "titulo": "Seu gargalo oculto está no SISTEMA DE GESTÃO",
        "diagnostico": ("Falta o leme estratégico. Sem planejamento, metas claras (indicador, prazo e valor) e tempo "
                        "dedicado a pensar o negócio, a empresa vive presa no operacional e apenas reage ao mercado. "
                        "O crescimento travou porque ninguém está, de fato, pilotando o longo prazo."),
    },
    {
        "id": "pessoas", "name": "Pessoas", "weight": 10, "color": "A855F7",
        "intro": "Time, liderança e cultura que sustentam o crescimento.",
        "questions": [
            "Existem ações de comunicação interna/endomarketing para gerar engajamento nas pessoas.",
            "O processo de recrutamento, seleção e integração é adequadamente conduzido.",
            "Os colaboradores recebem feedback periódico sobre seu desempenho.",
            "Talento e perfil dos colaboradores são mapeados com alguma ferramenta de Assessment (DISC, MBTI ou outros).",
            "É feito um planejamento e descrição de cada cargo.",
            "A liderança (incluindo os donos) tem uma comunicação clara, flexível e adequada.",
            "A liderança (incluindo os donos) realiza e oportuniza treinamentos técnicos e comportamentais.",
            "A liderança (incluindo os donos) executa estratégias de engajamento da equipe.",
            "A liderança (incluindo os donos) delega de forma efetiva.",
            "O ambiente físico da empresa está adequado às necessidades das atividades e dos funcionários.",
        ],
        "titulo": "Seu gargalo oculto está nas PESSOAS",
        "diagnostico": ("Seu time é seu maior ativo — e, hoje, seu maior gargalo. Sem uma liderança que delega, engaja "
                        "e desenvolve, tudo volta para o dono e a empresa não escala além da capacidade de uma única pessoa. "
                        "O faturamento parou porque você é, ao mesmo tempo, o motor e o limite do negócio."),
    },
]

# ---------- Estilos ----------
BG = "0B1120"
DARK = Font(color="EEF2FB", name="Calibri")
WHITE = "FFFFFF"
title_font = Font(size=20, bold=True, color="FFFFFF", name="Calibri")
sub_font = Font(size=11, italic=True, color="FFFFFF", name="Calibri")
dim_font = Font(size=13, bold=True, color="FFFFFF", name="Calibri")
hdr_font = Font(size=11, bold=True, color="FFFFFF", name="Calibri")
q_font = Font(size=11, color="1F2937", name="Calibri")
sub_total_font = Font(size=11, bold=True, color="111827", name="Calibri")

thin = Side(style="thin", color="D1D5DB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


wb = Workbook()

# ===================== ABA: Diagnóstico =====================
ws = wb.active
ws.title = "Diagnóstico"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 78
ws.column_dimensions["C"].width = 34
ws.column_dimensions["D"].width = 12

# Cabeçalho
ws.merge_cells("A1:D1")
c = ws["A1"]; c.value = "DIAGNÓSTICO DO GARGALO OCULTO"; c.font = title_font
c.fill = fill(BG); c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:D2")
c = ws["A2"]; c.value = "O mapa que revela por que seu faturamento parou de crescer"
c.font = sub_font; c.fill = fill("16203A"); c.alignment = center
ws.row_dimensions[2].height = 22

instr = [
    "COMO PREENCHER:",
    "1) Em cada item, abra o menu suspenso na coluna \"Sua avaliação\" e escolha uma das 3 opções.",
    "2) A pontuação é calculada automaticamente. Não precisa mexer na coluna \"Pontos\".",
    "3) Ao terminar, vá até a aba \"Resultado\" para ver o radar e o seu gargalo oculto.",
]
r = 3
for i, line in enumerate(instr):
    ws.merge_cells(f"A{r}:D{r}")
    c = ws[f"A{r}"]; c.value = line
    c.font = Font(bold=(i == 0), color="374151", size=10.5)
    c.alignment = left
    r += 1
r += 1  # linha em branco

subtotal_rows = {}

for dim in DIMENSIONS:
    # Cabeçalho da dimensão
    ws.merge_cells(f"A{r}:D{r}")
    c = ws[f"A{r}"]
    c.value = f"  {dim['name'].upper()}  —  cada item \"satisfatório\" vale {dim['weight']} pts  ·  {dim['intro']}"
    c.font = dim_font; c.fill = fill(dim["color"]); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 26
    r += 1

    # Cabeçalho da tabela
    ws[f"A{r}"].value = "#"
    ws[f"B{r}"].value = "Item avaliado"
    ws[f"C{r}"].value = "Sua avaliação"
    ws[f"D{r}"].value = "Pontos"
    for col in ("A", "B", "C", "D"):
        cell = ws[f"{col}{r}"]
        cell.font = hdr_font; cell.fill = fill("374151")
        cell.alignment = center if col in ("A", "C", "D") else left
        cell.border = border
    r += 1

    first_q = r
    for idx, q in enumerate(dim["questions"], start=1):
        ws[f"A{r}"].value = idx
        ws[f"A{r}"].alignment = center; ws[f"A{r}"].border = border; ws[f"A{r}"].font = q_font
        ws[f"B{r}"].value = q
        ws[f"B{r}"].alignment = left; ws[f"B{r}"].border = border; ws[f"B{r}"].font = q_font
        ws[f"B{r}"].fill = fill("FFFFFF")
        # Coluna de avaliação (dropdown)
        cell = ws[f"C{r}"]
        cell.value = ""
        cell.alignment = center; cell.border = border; cell.fill = fill("FEF3C7")
        cell.font = Font(bold=True, color="92400E", size=10.5)
        # Pontos (fórmula)
        w = dim["weight"]
        pts = ws[f"D{r}"]
        pts.value = (f'=IF($C{r}="{OPT_FULL}",{w},'
                     f'IF($C{r}="{OPT_PARTIAL}",{w/2 if (w/2)%1 else int(w/2)},0))')
        pts.alignment = center; pts.border = border; pts.font = q_font
        ws.row_dimensions[r].height = 30
        r += 1
    last_q = r - 1

    # Subtotal
    ws[f"A{r}"].fill = fill("E5E7EB"); ws[f"A{r}"].border = border
    c = ws[f"B{r}"]; c.value = f"SUBTOTAL — {dim['name']} (máximo 100)"
    c.font = sub_total_font; c.fill = fill("E5E7EB"); c.alignment = Alignment(horizontal="right", vertical="center"); c.border = border
    c = ws[f"C{r}"]
    c.value = f'=SUM(D{first_q}:D{last_q})&" / 100"'
    c.font = sub_total_font; c.fill = fill("E5E7EB"); c.alignment = center; c.border = border
    pts = ws[f"D{r}"]; pts.value = f"=SUM(D{first_q}:D{last_q})"
    pts.font = sub_total_font; pts.fill = fill("E5E7EB"); pts.alignment = center; pts.border = border
    subtotal_rows[dim["id"]] = r
    ws.row_dimensions[r].height = 24
    r += 2

# ===================== ABA: Listas (oculta) =====================
wl = wb.create_sheet("Listas")
wl["A1"] = OPT_FULL
wl["A2"] = OPT_PARTIAL
wl["A3"] = OPT_NONE
wl.sheet_state = "hidden"

# Validação de dados (menu suspenso) referenciando a aba Listas
dv = DataValidation(type="list", formula1="=Listas!$A$1:$A$3", allow_blank=True, showDropDown=False)
dv.prompt = "Escolha uma das 3 opções"
dv.promptTitle = "Avaliação"
ws.add_data_validation(dv)
# Aplica a validação a todas as células de avaliação (coluna C das perguntas)
for dim in DIMENSIONS:
    sub = subtotal_rows[dim["id"]]
    first = sub - len(dim["questions"])
    for rr in range(first, sub):
        dv.add(ws[f"C{rr}"])

# ===================== ABA: Textos (oculta) p/ lookup =====================
wt = wb.create_sheet("Textos")
wt["A1"] = "Dimensão"; wt["B1"] = "Título"; wt["C1"] = "Diagnóstico"
for i, dim in enumerate(DIMENSIONS, start=2):
    wt[f"A{i}"] = dim["name"]
    wt[f"B{i}"] = dim["titulo"]
    wt[f"C{i}"] = dim["diagnostico"]
# Faixas de maturidade
wt["E1"] = "Mín"; wt["F1"] = "Faixa"; wt["G1"] = "Resumo"
bands = [
    (0, "Sobrevivência", "A empresa funciona na base do esforço pessoal e da reação. Pequenos ajustes estruturais geram saltos rápidos."),
    (40, "Organização", "Você saiu do caos, mas um gargalo específico segura o crescimento do restante. Destravá-lo leva ao próximo patamar."),
    (70, "Escala", "A base está sólida. O gargalo é fino e específico — o ajuste que torna o crescimento previsível e sustentável."),
]
for i, (mn, lab, res) in enumerate(bands, start=2):
    wt[f"E{i}"] = mn; wt[f"F{i}"] = lab; wt[f"G{i}"] = res
wt.sheet_state = "hidden"

# ===================== ABA: Resultado =====================
wr = wb.create_sheet("Resultado")
wr.sheet_view.showGridLines = False
wr.column_dimensions["A"].width = 22
wr.column_dimensions["B"].width = 14
wr.column_dimensions["C"].width = 16
wr.column_dimensions["D"].width = 14
wr.column_dimensions["E"].width = 60

wr.merge_cells("A1:E1")
c = wr["A1"]; c.value = "RESULTADO DO DIAGNÓSTICO"; c.font = title_font
c.fill = fill(BG); c.alignment = center
wr.row_dimensions[1].height = 36

# Tabela do radar
hdr_row = 4
headers = ["Dimensão", "Atual (%)", "Desejado (%)", "Prazo"]
for j, h in enumerate(headers):
    cell = wr.cell(row=hdr_row, column=1 + j, value=h)
    cell.font = hdr_font; cell.fill = fill("374151"); cell.alignment = center; cell.border = border

first_data = hdr_row + 1
for i, dim in enumerate(DIMENSIONS):
    rr = first_data + i
    wr.cell(row=rr, column=1, value=dim["name"]).border = border
    wr.cell(row=rr, column=1).font = Font(bold=True)
    # Atual = subtotal da aba Diagnóstico
    a = wr.cell(row=rr, column=2, value=f"=Diagnóstico!D{subtotal_rows[dim['id']]}")
    a.alignment = center; a.border = border
    # Desejado (editável, padrão 100)
    d = wr.cell(row=rr, column=3, value=100)
    d.alignment = center; d.border = border; d.fill = fill("FEF3C7")
    # Prazo (editável)
    p = wr.cell(row=rr, column=4, value="")
    p.alignment = center; p.border = border; p.fill = fill("FEF3C7")
last_data = first_data + len(DIMENSIONS) - 1

# Gráfico radar
chart = RadarChart()
chart.type = "filled"
chart.style = 26
chart.title = "Radar das 4 dimensões — Atual x Desejado"
data = Reference(wr, min_col=2, max_col=3, min_row=hdr_row, max_row=last_data)
cats = Reference(wr, min_col=1, min_row=first_data, max_row=last_data)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 11
chart.width = 16
wr.add_chart(chart, "A12")

# Bloco do gargalo / maturidade (abaixo do gráfico)
base = 30
atual_range = f"B{first_data}:B{last_data}"
name_range = f"A{first_data}:A{last_data}"

def block(row, label, value_formula, label_fill="111A2E", value_fill="16203A", big=False):
    wr.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1)
    lc = wr.cell(row=row, column=1, value=label)
    lc.font = Font(bold=True, color="FFFFFF"); lc.fill = fill(label_fill); lc.alignment = Alignment(vertical="center", wrap_text=True)
    wr.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    vc = wr.cell(row=row, column=2, value=value_formula)
    vc.fill = fill(value_fill); vc.alignment = Alignment(vertical="center", wrap_text=True)
    vc.font = Font(bold=big, color="FFFFFF", size=13 if big else 11)
    wr.row_dimensions[row].height = 40 if big else 56

wr.merge_cells(start_row=base - 1, start_column=1, end_row=base - 1, end_column=5)
hc = wr.cell(row=base - 1, column=1, value="SUA LEITURA AUTOMÁTICA")
hc.font = Font(bold=True, size=12, color="FBBF24"); hc.fill = fill(BG); hc.alignment = center
wr.row_dimensions[base - 1].height = 26

gargalo_formula = f'=INDEX({name_range},MATCH(MIN({atual_range}),{atual_range},0))'
block(base, "Gargalo oculto", gargalo_formula, label_fill="7F1D1D", value_fill="991B1B", big=True)
block(base + 1, "Título",
      f'=IFERROR(VLOOKUP(B{base},Textos!$A$2:$C$5,2,FALSE),"")')
block(base + 2, "Diagnóstico",
      f'=IFERROR(VLOOKUP(B{base},Textos!$A$2:$C$5,3,FALSE),"")')
block(base + 3, "Maturidade geral",
      f'=ROUND(AVERAGE({atual_range}),0)&"%  —  "&IFERROR(VLOOKUP(ROUND(AVERAGE({atual_range}),0),Textos!$E$2:$G$4,2,TRUE),"")')
block(base + 4, "O que fazer agora",
      f'=IFERROR(VLOOKUP(ROUND(AVERAGE({atual_range}),0),Textos!$E$2:$G$4,3,TRUE),"")')

# CTA
cta_row = base + 6
wr.merge_cells(start_row=cta_row, start_column=1, end_row=cta_row, end_column=5)
c = wr.cell(row=cta_row, column=1,
            value="Agora você sabe ONDE travou. A próxima pergunta vale seu próximo salto: COMO destravar?")
c.font = Font(bold=True, size=12, color="20160A"); c.fill = fill("F59E0B")
c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
wr.row_dimensions[cta_row].height = 44

wb.active = wb.sheetnames.index("Diagnóstico")
out = "Diagnostico-do-Gargalo-Oculto.xlsx"
wb.save(out)
print("Planilha gerada:", out)
